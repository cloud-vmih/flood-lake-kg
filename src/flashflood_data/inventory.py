"""Deterministic, read-only inventory primitives for legacy dataset assets."""

from __future__ import annotations

import csv
import json
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from flashflood_data.catalog import sha256_file
from flashflood_data.io_atomic import atomic_target
from flashflood_data.models import AssetKind, AssetRecord, ValidationResult


@dataclass(frozen=True)
class InventoryRule:
    """An explicit classification rule for one legacy asset family."""

    source_id: str
    version: str
    kind: AssetKind
    glob: str
    media_type: str
    bundle_suffixes: tuple[str, ...]


def bundle_member(primary: Path, suffix: str) -> Path:
    """Return the expected bundle member for *suffix*."""
    if suffix == ".shp.xml":
        return Path(f"{primary}.xml")
    return primary.with_suffix(suffix)


def read_checksum_cache(path: Path) -> dict[str, dict[str, int | str]]:
    """Read a checksum cache, treating an absent or invalid cache as empty."""
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        entries = payload["entries"]
        if not isinstance(entries, dict):
            return {}
        return {
            str(key): value
            for key, value in entries.items()
            if isinstance(value, dict)
            and isinstance(value.get("size"), int)
            and isinstance(value.get("mtime_ns"), int)
            and isinstance(value.get("checksum"), str)
        }
    except (OSError, KeyError, TypeError, ValueError, json.JSONDecodeError):
        return {}


def write_json_atomic(path: Path, payload: object) -> None:
    """Publish stable JSON without exposing a partially written report."""
    body = json.dumps(payload, indent=2, sort_keys=True, ensure_ascii=False) + "\n"
    with atomic_target(path) as partial:
        partial.write_text(body, encoding="utf-8")


def cached_sha256(
    path: Path,
    *,
    dataset_root: Path,
    old_entries: dict[str, dict[str, int | str]],
    new_entries: dict[str, dict[str, int | str]],
    rehash: bool,
) -> str:
    """Hash a file unless its path, size, and nanosecond mtime match the cache."""
    relative = path.relative_to(dataset_root).as_posix()
    stat = path.stat()
    cached = old_entries.get(relative)
    if (
        not rehash
        and cached is not None
        and cached["size"] == stat.st_size
        and cached["mtime_ns"] == stat.st_mtime_ns
    ):
        checksum = str(cached["checksum"])
    else:
        checksum = sha256_file(path)
    new_entries[relative] = {
        "checksum": checksum,
        "mtime_ns": stat.st_mtime_ns,
        "size": stat.st_size,
    }
    return checksum


def compound_checksum(member_checksums: list[tuple[str, str]]) -> str:
    """Fingerprint a bundle from each present relative suffix and member digest."""
    if len(member_checksums) == 1:
        return member_checksums[0][1]
    digest = sha256()
    for suffix, checksum in sorted(member_checksums):
        digest.update(suffix.encode("utf-8"))
        digest.update(b"\0")
        digest.update(checksum.encode("ascii"))
        digest.update(b"\n")
    return digest.hexdigest()


def validate_known_format(
    path: Path, media_type: str, members: tuple[Path, ...]
) -> ValidationResult:
    """Perform a lightweight readability and signature check for a legacy asset."""
    checks: dict[str, bool] = {"readable": False, "known_format": False}
    metrics: dict[str, int | str] = {"bundle_members": len(members)}
    try:
        for member in members:
            with member.open("rb") as stream:
                stream.read(1)
        checks["readable"] = True

        if media_type in {"application/x-esri-shapefile", "image/tiff"}:
            with path.open("rb") as stream:
                signature = stream.read(4)
            if media_type == "application/x-esri-shapefile":
                checks["known_format"] = signature == b"\x00\x00\x27\x0a"
            else:
                checks["known_format"] = signature in {
                    b"II*\x00",
                    b"II+\x00",
                    b"MM\x00*",
                    b"MM\x00+",
                }
        elif media_type == "application/zip":
            checks["known_format"] = zipfile.is_zipfile(path)
        elif media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                worksheet = workbook.active
                rows = worksheet.iter_rows(values_only=True)
                header = next(rows, ())
                if header and str(header[0]).strip().casefold() == "stt":
                    data_rows = sum(
                        bool(row) and isinstance(row[0], int) and not isinstance(row[0], bool)
                        for row in rows
                    )
                else:
                    data_rows = sum(any(value is not None for value in row) for row in rows)
                metrics["workbook_data_rows"] = data_rows
            finally:
                workbook.close()
            checks["known_format"] = True
        elif media_type == "text/csv":
            with path.open("r", encoding="utf-8-sig", newline="") as stream:
                next(csv.reader(stream))
            checks["known_format"] = True
        elif media_type == "text/x-python":
            compile(path.read_text(encoding="utf-8"), path.name, "exec")
            checks["known_format"] = True
        else:
            checks["known_format"] = True
    except (OSError, UnicodeError, SyntaxError, StopIteration, ValueError, zipfile.BadZipFile):
        return ValidationResult(
            passed=False,
            checks=checks,
            metrics=metrics,
            messages=("legacy asset failed a readability or format check",),
        )
    return ValidationResult(
        passed=all(checks.values()),
        checks=checks,
        metrics=metrics,
    )


def deterministic_report(records: list[AssetRecord]) -> dict[str, object]:
    """Summarize inventory records without timestamps or machine-local arguments."""
    sources: dict[str, dict[str, int]] = {}
    status_counts: dict[str, int] = {}
    for record in records:
        summary = sources.setdefault(
            record.source_id,
            {"asset_count": 0, "duplicate_count": 0, "total_bytes": 0},
        )
        summary["asset_count"] += 1
        summary["total_bytes"] += record.size_bytes
        if record.duplicate_of_asset_id is not None:
            summary["duplicate_count"] += 1
        status_counts[record.status.value] = status_counts.get(record.status.value, 0) + 1
    return {
        "asset_count": len(records),
        "duplicate_count": sum(record.duplicate_of_asset_id is not None for record in records),
        "sources": dict(sorted(sources.items())),
        "status_counts": dict(sorted(status_counts.items())),
        "total_bytes": sum(record.size_bytes for record in records),
    }


def file_timestamp(path: Path) -> datetime:
    """Represent legacy acquisition time with the stable file modification time."""
    return datetime.fromtimestamp(path.stat().st_mtime, tz=UTC)
